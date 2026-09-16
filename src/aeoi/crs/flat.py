"""The flat input format: an Excel workbook (or CSV files) with four sheets.

Sheets and columns are the public contract of the toolkit; the Excel template is generated from
this module (``aeoi crs template``) so that the two never drift apart.

- ``ReportingFI``: two columns ``field`` / ``value`` (one reporting institution per workbook).
- ``Accounts``: one row per reportable account, holder data inline.
- ``ControllingPersons``: one row per controlling person, linked to the account by ``key``.
- ``Payments``: one row per payment, linked by ``key``.

Conventions: multi-value cells use ``;`` (``CH;DE``); TINs and INs are ``value@CC`` per entry,
``@CC`` optional; booleans accept true/false, 1/0, yes/no, ja/nein, oui/non, si/no; dates are
ISO ``YYYY-MM-DD`` or real Excel dates; amounts use a dot as decimal separator or numeric cells.
"""

from __future__ import annotations

import csv
import datetime as dt
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from pydantic import ValidationError

from aeoi.crs import codes
from aeoi.crs.model import (
    Account,
    Address,
    ControllingPerson,
    Message,
    Organisation,
    Payment,
    Person,
    ReportingFI,
    Tin,
)

SEP = ";"
TRUE = {"true", "1", "yes", "y", "ja", "oui", "si", "sì", "wahr", "x"}
FALSE = {"false", "0", "no", "n", "nein", "non", "falsch", ""}


@dataclass(frozen=True)
class Column:
    name: str
    description: str
    codes: dict | None = None  # allowed values -> meaning (for validation lists)
    required: bool = False
    kind: str = "text"  # text | bool | date | decimal | int | multi | tins


ADDRESS_COLUMNS = [
    Column(
        "address_country", "Country of the address, ISO 3166-1 alpha-2 (e.g. CH)", required=True
    ),
    Column("address_street", "Street (AddressFix.Street)"),
    Column("address_building", "Building identifier / house number"),
    Column("address_suite", "Suite identifier"),
    Column("address_floor", "Floor identifier"),
    Column("address_district", "District name"),
    Column("address_pob", "P.O. box"),
    Column("address_post_code", "Postal code"),
    Column("address_city", "City - mandatory, AddressFix must be used (ESTV 98104)", required=True),
    Column("address_subentity", "Country subentity (canton, state)"),
    Column("address_free", "Optional free-text address in addition to the fixed parts"),
    Column("legal_address_type", "OECD301-305", codes.LEGAL_ADDRESS_TYPE),
]

PERSON_COLUMNS = [
    Column("first_name", "First name; NFN if none (ESTV 5.3.8)", required=True),
    Column("last_name", "Last name", required=True),
    Column("middle_name", "Middle name"),
    Column("name_type", "OECD202-208 (OECD201 not allowed)", codes.NAME_TYPE),
    Column("birth_date", "Date of birth YYYY-MM-DD (after 1900-01-01)", kind="date"),
    Column("birth_city", "City of birth"),
    Column("birth_country", "Country of birth, ISO alpha-2"),
    Column("residence_countries", "Tax residence countries, ISO alpha-2, ';'-separated", required=True, kind="multi"),
    Column("tins", "TINs as value@CC;value@CC (CC = issuing country, optional)", kind="tins"),
    Column("nationalities", "Nationalities, ISO alpha-2, ';'-separated", kind="multi"),
]  # fmt: skip

REPORTING_FI_FIELDS = [
    Column("estv_id", "ESTV-ID of the reporting FI (SendingCompanyIN), e.g. 052.0000.0000", required=True),
    Column("uid", "UID of the reporting FI (IN), e.g. CHE-123.456.789", required=True),
    Column("name", "Legal name of the reporting FI", required=True),
    Column("contact", "Contact for queries (MessageSpec.Contact), optional"),
    Column("reporting_year", "Reporting year (calendar year the data refers to)", required=True, kind="int"),
    Column("message_type_indic", "CRS701 new data / CRS703 nil report (no accounts)", codes.MESSAGE_TYPE_INDIC),
    *ADDRESS_COLUMNS,
]  # fmt: skip

ACCOUNT_COLUMNS = [
    Column("key", "Row identifier, unique per workbook; links ControllingPersons and Payments", required=True),
    Column("doc_ref_id", "Leave empty: generated (CH<year>CH<uuid>). Fill only to reuse a known DocRefId"),
    Column("account_number", "Account number; NANUM if none (ESTV 5.3.7)", required=True),
    Column("account_number_type", "OECD601 IBAN, 602 OBAN, 603 ISIN, 604 OSIN, 605 Other, 606 e-money", codes.ACCT_NUMBER_TYPE),
    Column("undocumented", "true/false - undocumented account (holder must be CH individual)", kind="bool"),
    Column("closed", "true/false - closed during the year (balance must be 0)", kind="bool"),
    Column("dormant", "true/false - dormant account", kind="bool"),
    Column("holder_type", "individual or organisation", {"individual": "natural person", "organisation": "entity"}, required=True),
    *PERSON_COLUMNS,
    Column("org_name", "Organisation: legal name"),
    Column("org_name_type", "Organisation: OECD202-208", codes.NAME_TYPE),
    Column("acct_holder_type", "Organisation: CRS101 passive NFE with CPs / CRS102 reportable person / CRS103 passive NFE reportable", codes.ACCT_HOLDER_TYPE),
    Column("org_ins", "Organisation: INs as value@CC;value@CC", kind="tins"),
    *ADDRESS_COLUMNS,
    Column("balance", "Account balance at year end, >= 0 (0 if closed)", required=True, kind="decimal"),
    Column("currency", "ISO 4217 currency of the balance (CHF, EUR, USD ...)", required=True),
    Column("self_cert", "3.0: CRS901 valid self-certification / CRS902 none", codes.SELF_CERT),
    Column("dd_procedure", "3.0: CRS1201 new account / CRS1202 preexisting account", codes.DD_PROCEDURE),
    Column("account_type", "3.0: CRS1101 depository / 1102 custodial / 1103 insurance-annuity / 1104 equity-debt interest", codes.ACCOUNT_TYPE),
    Column("joint_account_number", "3.0: number of joint holders (1-200), optional", kind="int"),
    Column("equity_interest_types", "3.0: CRS401-410, ';'-separated, only with account_type CRS1104", codes.EQUITY_INTEREST_TYPE, kind="multi"),
]  # fmt: skip

CONTROLLING_PERSON_COLUMNS = [
    Column("key", "Key of the account (Accounts.key)", required=True),
    *PERSON_COLUMNS,
    *ADDRESS_COLUMNS,
    Column("ctrlg_person_types", "CRS801-813, ';'-separated (3.0: at least one)", codes.CTRLG_PERSON_TYPE, kind="multi"),
    Column("self_cert", "3.0: CRS1001 valid self-certification / CRS1002 none", codes.SELF_CERT_CP),
]  # fmt: skip

PAYMENT_COLUMNS = [
    Column("key", "Key of the account (Accounts.key)", required=True),
    Column("payment_type", "CRS501 dividends / 502 interest / 503 gross proceeds / 504 other", codes.PAYMENT_TYPE, required=True),
    Column("amount", "Amount, >= 0", required=True, kind="decimal"),
    Column("currency", "ISO 4217 currency", required=True),
]  # fmt: skip

SHEETS = {
    "ReportingFI": REPORTING_FI_FIELDS,
    "Accounts": ACCOUNT_COLUMNS,
    "ControllingPersons": CONTROLLING_PERSON_COLUMNS,
    "Payments": PAYMENT_COLUMNS,
}


@dataclass(frozen=True)
class InputProblem:
    sheet: str
    row: int | None
    column: str | None
    message: str

    def __str__(self) -> str:
        loc = self.sheet
        if self.row is not None:
            loc += f" row {self.row}"
        if self.column:
            loc += f", column {self.column}"
        return f"{loc}: {self.message}"


@dataclass
class ReadResult:
    message: Message | None
    problems: list[InputProblem] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return self.message is not None and not self.problems


# --- cell parsing ----------------------------------------------------------------------------


def _text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _bool(v: Any) -> bool | None:
    if isinstance(v, bool):
        return v
    s = _text(v).lower()
    if s in TRUE:
        return True
    if s in FALSE:
        return False
    return None


def _date(v: Any) -> dt.date | None | str:
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = _text(v)
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(s, fmt).replace(tzinfo=dt.UTC).date()
        except ValueError:
            continue
    return "invalid"


def _decimal(v: Any) -> Decimal | None:
    if v is None or v == "":
        return None
    try:
        return Decimal(str(v).replace("'", "").replace(" ", "").replace(",", "."))
    except InvalidOperation:
        return None


def _multi(v: Any) -> list[str]:
    return [x.strip().upper() for x in _text(v).split(SEP) if x.strip()]


def _tins(v: Any) -> list[Tin]:
    out = []
    for item in [x.strip() for x in _text(v).split(SEP) if x.strip()]:
        value, _, cc = item.partition("@")
        out.append(Tin(value=value.strip(), issued_by=cc.strip().upper() or None))
    return out


def _address(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "country": _text(row.get("address_country")).upper(),
        "city": _text(row.get("address_city")),
        "street": _text(row.get("address_street")) or None,
        "building_identifier": _text(row.get("address_building")) or None,
        "suite_identifier": _text(row.get("address_suite")) or None,
        "floor_identifier": _text(row.get("address_floor")) or None,
        "district_name": _text(row.get("address_district")) or None,
        "pob": _text(row.get("address_pob")) or None,
        "post_code": _text(row.get("address_post_code")) or None,
        "country_subentity": _text(row.get("address_subentity")) or None,
        "address_free": _text(row.get("address_free")) or None,
        "legal_address_type": _text(row.get("legal_address_type")).upper() or None,
    }


def _person(row: dict[str, Any], problems: list[InputProblem], sheet: str, rownum: int) -> dict:
    birth = _date(row.get("birth_date"))
    if birth == "invalid":
        problems.append(InputProblem(sheet, rownum, "birth_date", "not a date (use YYYY-MM-DD)"))
        birth = None
    return {
        "first_name": _text(row.get("first_name")),
        "last_name": _text(row.get("last_name")),
        "middle_name": _text(row.get("middle_name")) or None,
        "name_type": _text(row.get("name_type")).upper() or None,
        "birth_date": birth,
        "birth_city": _text(row.get("birth_city")) or None,
        "birth_country": _text(row.get("birth_country")).upper() or None,
        "residence_countries": _multi(row.get("residence_countries")),
        "tins": _tins(row.get("tins")),
        "nationalities": _multi(row.get("nationalities")),
        "address": _address(row),
    }


# --- workbook / csv access -------------------------------------------------------------------


def _rows_from_xlsx(path: Path) -> dict[str, list[dict[str, Any]]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out: dict[str, list[dict[str, Any]]] = {}
    for name in SHEETS:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            out[name] = []
            continue
        if name == "ReportingFI":
            out[name] = [
                {"field": _text(r[0]), "value": r[1] if len(r) > 1 else None, "_row": i + 1}
                for i, r in enumerate(rows)
                if r and _text(r[0]) and _text(r[0]) != "field"
            ]
            continue
        header = [_text(h) for h in rows[0]]
        out[name] = [
            {**dict(zip(header, r, strict=False)), "_row": i + 2}
            for i, r in enumerate(rows[1:])
            if r and any(_text(c) for c in r)
        ]
    return out


def _rows_from_csv_dir(folder: Path) -> dict[str, list[dict[str, Any]]]:
    out: dict[str, list[dict[str, Any]]] = {}
    for name in SHEETS:
        f = folder / f"{name}.csv"
        if not f.exists():
            continue
        with f.open(encoding="utf-8-sig", newline="") as fh:
            reader = csv.DictReader(fh, delimiter=";") if name != "ReportingFI" else None
            if name == "ReportingFI":
                fh.seek(0)
                out[name] = [
                    {"field": r[0].strip(), "value": r[1] if len(r) > 1 else None, "_row": i + 1}
                    for i, r in enumerate(csv.reader(fh, delimiter=";"))
                    if r and r[0].strip() and r[0].strip() != "field"
                ]
            else:
                out[name] = [{**r, "_row": i + 2} for i, r in enumerate(reader)]
    return out


def read(path: str | Path) -> ReadResult:
    """Read a workbook (``.xlsx``) or a folder of ``<Sheet>.csv`` files into a :class:`Message`."""
    path = Path(path)
    problems: list[InputProblem] = []
    sheets = _rows_from_csv_dir(path) if path.is_dir() else _rows_from_xlsx(path)
    for name in ("ReportingFI", "Accounts"):
        if name not in sheets:
            problems.append(InputProblem(name, None, None, "sheet is missing"))
    if problems:
        return ReadResult(None, problems)

    fi_values = {r["field"]: r["value"] for r in sheets["ReportingFI"]}
    unknown = set(fi_values) - {c.name for c in REPORTING_FI_FIELDS}
    for u in sorted(unknown):
        problems.append(InputProblem("ReportingFI", None, u, "unknown field"))
    year = _text(fi_values.get("reporting_year"))
    try:
        reporting_year = int(year)
    except ValueError:
        problems.append(InputProblem("ReportingFI", None, "reporting_year", "must be a year"))
        reporting_year = 0
    try:
        fi = ReportingFI(
            estv_id=_text(fi_values.get("estv_id")),
            uid=_text(fi_values.get("uid")).upper(),
            name=_text(fi_values.get("name")),
            contact=_text(fi_values.get("contact")) or None,
            address=Address(**_address({k: v for k, v in fi_values.items()})),
        )
    except ValidationError as exc:
        for e in exc.errors():
            problems.append(
                InputProblem("ReportingFI", None, ".".join(map(str, e["loc"])), e["msg"])
            )
        return ReadResult(None, problems)

    cps: dict[str, list[ControllingPerson]] = {}
    for r in sheets.get("ControllingPersons", []):
        key = _text(r.get("key"))
        if not key:
            problems.append(
                InputProblem("ControllingPersons", r["_row"], "key", "missing account key")
            )
            continue
        try:
            cp = ControllingPerson(
                person=Person(**_person(r, problems, "ControllingPersons", r["_row"])),
                ctrlg_person_types=_multi(r.get("ctrlg_person_types")),
                self_cert=_text(r.get("self_cert")).upper() or None,
            )
        except ValidationError as exc:
            for e in exc.errors():
                problems.append(
                    InputProblem(
                        "ControllingPersons", r["_row"], ".".join(map(str, e["loc"])), e["msg"]
                    )
                )
            continue
        cps.setdefault(key, []).append(cp)

    pays: dict[str, list[Payment]] = {}
    for r in sheets.get("Payments", []):
        key = _text(r.get("key"))
        amount = _decimal(r.get("amount"))
        if not key:
            problems.append(InputProblem("Payments", r["_row"], "key", "missing account key"))
            continue
        if amount is None:
            problems.append(InputProblem("Payments", r["_row"], "amount", "not a number"))
            continue
        pays.setdefault(key, []).append(
            Payment(
                payment_type=_text(r.get("payment_type")).upper(),
                amount=amount,
                currency=_text(r.get("currency")).upper(),
            )
        )

    accounts: list[Account] = []
    seen_keys: set[str] = set()
    for r in sheets["Accounts"]:
        rownum = r["_row"]
        key = _text(r.get("key"))
        if not key:
            problems.append(InputProblem("Accounts", rownum, "key", "missing key"))
            continue
        seen_keys.add(key)
        holder_type = _text(r.get("holder_type")).lower()
        person = organisation = None
        if holder_type == "individual":
            person = _person(r, problems, "Accounts", rownum)
        elif holder_type == "organisation":
            organisation = {
                "name": _text(r.get("org_name")),
                "name_type": _text(r.get("org_name_type")).upper() or None,
                "acct_holder_type": _text(r.get("acct_holder_type")).upper(),
                "residence_countries": _multi(r.get("residence_countries")),
                "ins": _tins(r.get("org_ins")),
                "address": _address(r),
            }
        else:
            problems.append(
                InputProblem(
                    "Accounts", rownum, "holder_type", "must be individual or organisation"
                )
            )
            continue
        balance = _decimal(r.get("balance"))
        row_ok = balance is not None
        if balance is None:
            problems.append(InputProblem("Accounts", rownum, "balance", "not a number"))
        flags = {}
        for name in ("undocumented", "closed", "dormant"):
            b = _bool(r.get(name))
            if b is None:
                problems.append(InputProblem("Accounts", rownum, name, "must be true or false"))
                b = False
            flags[name] = b
        if not row_ok:
            continue
        joint = _text(r.get("joint_account_number"))
        try:
            acc = Account(
                key=key,
                doc_ref_id=_text(r.get("doc_ref_id")) or None,
                account_number=_text(r.get("account_number")),
                account_number_type=_text(r.get("account_number_type")).upper() or None,
                holder_person=Person(**person) if person else None,
                holder_organisation=Organisation(**organisation) if organisation else None,
                controlling_persons=cps.pop(key, []),
                balance=balance,
                currency=_text(r.get("currency")).upper(),
                payments=pays.pop(key, []),
                self_cert=_text(r.get("self_cert")).upper() or None,
                dd_procedure=_text(r.get("dd_procedure")).upper() or None,
                account_type=_text(r.get("account_type")).upper() or None,
                joint_account_number=int(joint) if joint else None,
                equity_interest_types=_multi(r.get("equity_interest_types")),
                **flags,
            )
        except (ValidationError, ValueError) as exc:
            if isinstance(exc, ValidationError):
                for e in exc.errors():
                    problems.append(
                        InputProblem("Accounts", rownum, ".".join(map(str, e["loc"])), e["msg"])
                    )
            else:
                problems.append(InputProblem("Accounts", rownum, "joint_account_number", str(exc)))
            continue
        accounts.append(acc)

    for key in cps:
        problems.append(
            InputProblem("ControllingPersons", None, "key", f"no account with key {key!r}")
        )
    for key in pays:
        problems.append(InputProblem("Payments", None, "key", f"no account with key {key!r}"))

    msg = Message(
        reporting_fi=fi,
        reporting_year=reporting_year,
        accounts=accounts,
        message_type_indic=_text(fi_values.get("message_type_indic")).upper() or "CRS701",
    )
    return ReadResult(msg, problems)


_NAME_RE = re.compile(r"^[a-z_]+$")
assert all(_NAME_RE.match(c.name) for cols in SHEETS.values() for c in cols)
