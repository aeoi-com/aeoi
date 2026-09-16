"""The CLI chain: template --example -> check -> build (2.0 and 3.0) -> package -> inspect."""

from pathlib import Path

import pytest
import xmlschema
from cryptography.hazmat.primitives import serialization
from cryptography.hazmat.primitives.asymmetric import rsa

from aeoi import cli

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def key_pem(tmp_path_factory) -> Path:
    private = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    path = tmp_path_factory.mktemp("keys") / "placeholder-public.pem"
    path.write_bytes(
        private.public_key().public_bytes(
            serialization.Encoding.PEM, serialization.PublicFormat.SubjectPublicKeyInfo
        )
    )
    return path


def test_cli_chain(tmp_path, key_pem, capsys):
    xlsx = tmp_path / "example.xlsx"
    assert cli.main(["crs", "template", "--out", str(xlsx), "--example"]) == 0
    assert cli.main(["crs", "check", "--input", str(xlsx), "--version", "3.0"]) == 0
    assert cli.main(["crs", "check", "--input", str(xlsx), "--version", "2.0"]) == 0
    for version in ("2.0", "3.0"):
        xml = tmp_path / f"report-{version}.xml"
        pkg = tmp_path / f"Test-report-{version}.zip"
        rc = cli.main(
            ["crs", "build", "--input", str(xlsx), "--version", version, "--out", str(xml),
             "--test", "--key", str(key_pem), "--package", str(pkg)]
        )  # fmt: skip
        assert rc == 0
        schema = xmlschema.XMLSchema(
            ROOT / "schemas" / f"crs_v{version}" / f"CrsXML_v{version}.xsd"
        )
        assert schema.is_valid(xml)
        assert cli.main(["estv", "inspect", str(pkg), "--key", str(key_pem), "--test"]) == 0
    out = capsys.readouterr().out
    assert '"package_sha256"' in out


def test_cli_reports_problems_and_fails(tmp_path, capsys):
    xlsx = tmp_path / "example.xlsx"
    cli.main(["crs", "template", "--out", str(xlsx), "--example"])
    import openpyxl

    wb = openpyxl.load_workbook(xlsx)
    ws = wb["Accounts"]
    header = [c.value for c in ws[1]]
    ws.cell(row=2, column=header.index("self_cert") + 1).value = None
    wb.save(xlsx)
    assert cli.main(["crs", "check", "--input", str(xlsx), "--version", "3.0"]) == 1
    err = capsys.readouterr().err
    assert "Accounts[key=A1].self_cert" in err and "[3.0]" in err
    # the same file is fine for 2.0
    assert cli.main(["crs", "check", "--input", str(xlsx), "--version", "2.0"]) == 0
