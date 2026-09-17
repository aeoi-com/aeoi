"""The Excel template speaks German, French and Italian: every text it writes has a translation,
and a template written in a language carries that language in its comments, ReadMe and Codes."""

from __future__ import annotations

import openpyxl

from aeoi.crs import codes, flat, template

CODE_TABLES = (
    "ACCT_HOLDER_TYPE",
    "EQUITY_INTEREST_TYPE",
    "PAYMENT_TYPE",
    "MESSAGE_TYPE_INDIC",
    "CTRLG_PERSON_TYPE",
    "SELF_CERT",
    "SELF_CERT_CP",
    "ACCOUNT_TYPE",
    "DD_PROCEDURE",
    "TRANSITIONAL",
    "NAME_TYPE",
    "LEGAL_ADDRESS_TYPE",
    "ACCT_NUMBER_TYPE",
)


def template_strings() -> set[str]:
    strings = {c.description for cols in flat.SHEETS.values() for c in cols}
    for name in CODE_TABLES:
        strings |= set(getattr(codes, name).values())
    strings |= set(template.STEPS)
    strings |= {template.TITLE, template.SUMMARY, template.CLOSING, "Five steps"}
    strings |= {"field", "value", "description", "code list", "code", "meaning"}
    strings.add("mandatory when holder_type = {value}")
    for heading, paragraphs in template.SECTIONS:
        strings |= {heading, *paragraphs}
    return strings


def test_every_template_text_has_every_language():
    tr = template.translations()
    missing = [
        (lang, s)
        for s in sorted(template_strings())
        for lang in ("de", "fr", "it")
        if not tr.get(s, {}).get(lang)
    ]
    assert not missing, missing


def test_placeholders_survive_translation():
    for text, entry in template.translations().items():
        for lang, value in entry.items():
            assert ("{value}" in text) == ("{value}" in value), (lang, text)


def test_template_in_german(tmp_path):
    path = tmp_path / "vorlage-de.xlsx"
    template.write_template(path, lang="de")
    wb = openpyxl.load_workbook(path)
    accounts = wb["Accounts"]
    comments = {c.value: c.comment.text for c in accounts[1] if c.comment}
    assert comments["account_number"].startswith("Kontonummer")
    assert "Pflicht, wenn holder_type = individual" in comments["first_name"]
    assert wb["ReadMe"]["A1"].value == "meldbar - CRS-Vorlage für AIA-Meldungen"
    assert wb["ReportingFI"]["A1"].value == "Feld"
    meanings = {row[1]: row[2] for row in wb["Codes"].iter_rows(min_row=2, values_only=True)}
    assert meanings["CRS1101"] == "Einlagenkonto"
    assert meanings["CRS703"].startswith("Die Meldung teilt mit")


def test_template_english_is_unchanged_and_readable(tmp_path):
    path = tmp_path / "template.xlsx"
    template.write_template(path)
    wb = openpyxl.load_workbook(path)
    assert wb["ReportingFI"]["A1"].value == "field"
    result = flat.read(path)  # an empty template reads without crashing and reports what is missing
    assert any(p.column == "reporting_year" for p in result.problems)
