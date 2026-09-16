"""Flat input: template -> filled workbook -> read -> domain -> XML, plus talking input errors."""

import datetime as dt
from pathlib import Path

import openpyxl
import pytest
import xmlschema

from aeoi.crs import build, flat, model, template
from aeoi.crs.example import sample_message

ROOT = Path(__file__).resolve().parents[1]
TODAY = dt.date(2027, 3, 1)


@pytest.fixture(scope="module")
def schema_v3():
    return xmlschema.XMLSchema(ROOT / "schemas" / "crs_v3.0" / "CrsXML_v3.0.xsd")


def test_template_has_all_sheets_and_headers(tmp_path):
    path = tmp_path / "template.xlsx"
    template.write_template(path)
    wb = openpyxl.load_workbook(path)
    assert set(flat.SHEETS) <= set(wb.sheetnames)
    assert "Codes" in wb.sheetnames and "ReadMe" in wb.sheetnames
    for name, columns in flat.SHEETS.items():
        ws = wb[name]
        if name == "ReportingFI":
            fields = [ws.cell(row=i, column=1).value for i in range(2, len(columns) + 2)]
            assert fields == [c.name for c in columns]
        else:
            header = [c.value for c in ws[1]]
            assert header == [c.name for c in columns]
            assert ws.data_validations.dataValidation, name  # drop-downs exist


def test_round_trip_workbook_to_valid_xml(tmp_path, schema_v3):
    original = sample_message()
    path = tmp_path / "filled.xlsx"
    template.write_message(original, path)
    result = flat.read(path)
    assert result.ok, [str(p) for p in result.problems]
    msg = result.message
    assert msg.reporting_fi == original.reporting_fi
    assert msg.reporting_year == original.reporting_year
    assert msg.accounts == original.accounts
    rep = model.check_message(msg, "3.0", today=TODAY)
    assert not [p for p in rep.problems if p.rule != "info"], rep.problems
    xml = build.build(msg, "3.0").xml
    assert schema_v3.is_valid(xml)


def test_csv_folder_is_accepted(tmp_path):
    fi = tmp_path / "ReportingFI.csv"
    fi.write_text(
        "field;value\nestv_id;052.0000.0000\nuid;CHE-123.456.789\nname;Beispiel AG\n"
        "reporting_year;2026\naddress_country;CH\naddress_city;Zürich\n",
        encoding="utf-8",
    )
    (tmp_path / "Accounts.csv").write_text(
        "key;account_number;holder_type;first_name;last_name;residence_countries;"
        "address_country;address_city;balance;currency;closed;undocumented;dormant\n"
        "A1;NANUM;individual;NFN;Rossi;IT;IT;Roma;10.5;EUR;false;false;false\n",
        encoding="utf-8",
    )
    result = flat.read(tmp_path)
    assert result.ok, [str(p) for p in result.problems]
    assert result.message.accounts[0].holder_person.last_name == "Rossi"


def test_input_problems_name_sheet_row_and_column(tmp_path):
    msg = sample_message()
    path = tmp_path / "bad.xlsx"
    template.write_message(msg, path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Accounts"]
    header = [c.value for c in ws[1]]
    ws.cell(row=2, column=header.index("balance") + 1, value="abc")
    ws.cell(row=3, column=header.index("holder_type") + 1, value="trust")
    ws.cell(row=2, column=header.index("closed") + 1, value="maybe")
    wb["Payments"].cell(row=2, column=1, value="A9")  # payment for an unknown account
    wb.save(path)
    result = flat.read(path)
    texts = [str(p) for p in result.problems]
    assert any(t.startswith("Accounts row 2, column balance") for t in texts), texts
    assert any(t.startswith("Accounts row 3, column holder_type") for t in texts), texts
    assert any("column closed: must be true or false" in t for t in texts), texts
    assert any("no account with key 'A9'" in t for t in texts), texts


def test_booleans_dates_and_multivalues():
    assert flat._bool("Ja") is True and flat._bool("nein") is False and flat._bool("?") is None
    assert flat._date("12.04.1975") == dt.date(1975, 4, 12)
    assert flat._date("1975-04-12") == dt.date(1975, 4, 12)
    assert flat._date("not a date") == "invalid"
    assert flat._multi(" ch; de ;") == ["CH", "DE"]
    tins = flat._tins("123@de; 456")
    assert (tins[0].value, tins[0].issued_by, tins[1].issued_by) == ("123", "DE", None)
