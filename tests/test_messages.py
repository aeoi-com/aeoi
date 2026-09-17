"""The message catalogue: every id has every language with the same placeholders, every Msg the
code builds renders in German, reports and workflow errors follow the requested language."""

from __future__ import annotations

import pickle
import re
from pathlib import Path

import pytest

from aeoi import messages
from aeoi.crs import build, check, template, validate, workflow
from aeoi.crs.example import sample_message
from aeoi.messages import Msg
from aeoi.registry import Registry

SRC = Path(__file__).resolve().parents[1] / "src" / "aeoi"
PLACEHOLDER = re.compile(r"\{(\w+)\}")


def test_catalogue_complete_and_consistent():
    cat = messages.catalogue()
    assert len(cat) > 120
    for msg_id, entry in cat.items():
        for lang in messages.LANGUAGES:
            assert entry.get(lang), (msg_id, lang)
        fields = {lang: set(PLACEHOLDER.findall(entry[lang])) for lang in messages.LANGUAGES}
        assert len(set(map(frozenset, fields.values()))) == 1, (msg_id, fields)


def test_every_msg_id_in_the_code_exists():
    used: set[str] = set()
    for path in SRC.rglob("*.py"):
        for a, b in re.findall(
            r'Msg\(\s*"([a-z0-9_]+)"(?:\s+if\s+[^"]+?\s+else\s+"([a-z0-9_]+)")?',
            path.read_text(encoding="utf-8"),
        ):
            used |= {a, b} - {""}
    used |= {
        "st_built",
        "st_submitted",
        "st_accepted",
        "st_rejected",
        "st_discarded",
    }  # built dynamically
    used.discard("st_")  # the prefix of the dynamic status ids
    missing = sorted(used - set(messages.catalogue()))
    assert not missing, missing
    unused = sorted(set(messages.catalogue()) - used)
    assert not unused, unused


def test_msg_is_a_string_and_renders_other_languages():
    m = Msg("code_invalid", value="'X'", options="A, B")
    assert m == "'X' is not one of A, B" and isinstance(m, str)
    assert m.text("de") == "'X' ist keiner der Werte A, B"
    assert m.text("xx") == str(m)  # unknown language: English
    nested = Msg("charset", text="'#'", position=9, reason=Msg("char_excluded", code="0023"))
    assert nested.text("de") == "'#' an Position 9: durch Anhang 7.2 ausgeschlossen (U+0023)"
    listed = Msg("wf_blocked", problems=[Msg("must_be_ch"), Msg("must_be_crs")])
    assert (
        listed == "must be CH; must be CRS" and listed.text("de") == "muss CH sein; muss CRS sein"
    )
    assert (
        Msg("partner_list_missing", year=2026).text("de") == "keine Partnerstaaten-Liste für 2026"
    )
    copy = pickle.loads(pickle.dumps(nested))
    assert copy == nested and copy.text("de") == nested.text("de")


def test_reports_render_in_german(tmp_path):
    xml = build.build(sample_message(), "3.0", test=True).xml
    bad = tmp_path / "Test-bad.xml"
    bad.write_text(
        xml.replace("Beispiel AG", "Beispiel # AG").replace(
            "<crs:ResCountryCode>DE", "<crs:ResCountryCode>US", 1
        ),
        encoding="utf-8",
    )
    rep = validate.validate_file(bad)
    de = rep.render("de")
    assert "an Position 9: durch Anhang 7.2 ausgeschlossen (U+0023) [50005]" in de
    assert "keines der Länder US war 2026 ein AIA-Partnerstaat der Schweiz" in de
    assert "die USA sind kein AIA-Partnerstaat" in de
    en = rep.render()
    assert "at position 9: excluded by Anhang 7.2 (U+0023) [50005]" in en
    assert de.split("\n")[0] == en.split("\n")[0]  # the headline is language-neutral
    for p in (
        rep.problems + rep.infos
    ):  # every message the engine produced renders without leftovers
        assert "{" not in messages.text(p.message, "de")

    book = tmp_path / "book.xlsx"
    template.write_message(sample_message(), book)
    from openpyxl import load_workbook

    wb = load_workbook(book)
    ws = wb["Accounts"]
    ws.cell(row=2, column=[c.value for c in ws[1]].index("balance") + 1).value = "abc"
    wb.save(book)
    report = check.check_workbook(book, "3.0")
    assert "Accounts row 2, column balance: keine Zahl" in report.render("de")
    assert "Accounts row 2, column balance: not a number" in report.render()


def test_workflow_errors_render_in_german(tmp_path):
    with Registry(tmp_path / "r.sqlite") as reg:
        workflow.build(sample_message(), "3.0", reg, test=True)
        intent = workflow.plan(sample_message(), "3.0", reg, test=True)
        problems = intent.as_dict("de")["problems"]
        assert (
            problems
            and "nicht angenommen; vor der Korrektur das Ergebnis des Portals erfassen"
            in problems[0]
        )
        assert "ist erstellt, Ergebnis offen" in problems[0]
        with pytest.raises(workflow.WorkflowError) as exc:
            workflow.build(sample_message(), "3.0", reg, test=True)
        assert messages.text(exc.value.args[0], "de").startswith("Konto 'A1': Meldung")
        assert str(exc.value).startswith("account 'A1': message")
