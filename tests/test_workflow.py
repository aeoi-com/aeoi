"""The one-call operations the browser page uses: plan, build (new / correction / both), encrypt,
record the portal outcome, remember the ESTV key. The registry is the state."""

from __future__ import annotations

import pytest
from cryptography.hazmat.primitives import serialization
from cryptography.hazmat.primitives.asymmetric import rsa

from aeoi.crs import template, workflow
from aeoi.crs.example import sample_message
from aeoi.estv import packaging
from aeoi.registry import Registry


@pytest.fixture
def reg(tmp_path):
    with Registry(tmp_path / "institut.sqlite") as r:
        yield r


@pytest.fixture(scope="module")
def key_pair():
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    pem = key.public_key().public_bytes(
        serialization.Encoding.PEM, serialization.PublicFormat.SubjectPublicKeyInfo
    )
    return key, pem


def _indics(xml: str) -> list[str]:
    import re

    return re.findall(r"<stf:DocTypeIndic>(OECD1\d)</stf:DocTypeIndic>", xml)


def _accept(reg, ref):
    workflow.record_outcome(reg, f"MessageRefId {ref}\nStatus: accepted\n", message_ref_id=ref)


def test_plan_without_registry_only_allows_tests():
    msg = sample_message()
    intent = workflow.plan(msg, "3.0", None, test=True)
    assert intent.messages == ["new"] and intent.new_keys == ["A1", "A2"] and not intent.problems
    prod = workflow.plan(msg, "3.0", None, test=False)
    assert prod.problems and "registry" in prod.problems[0]
    with pytest.raises(workflow.WorkflowError, match="registry"):
        workflow.build(msg, "3.0", None, test=False)


def test_new_then_correction_then_readd(reg, key_pair):
    key, pem = key_pair
    msg = sample_message()
    # 1. first message: everything new, encrypted with the remembered key
    assert workflow.remember_public_key(reg, pem)
    view = workflow.registry_view(reg)
    assert view["open"] and view["has_key"] and view["messages"] == []
    outs = workflow.build(msg, "3.0", reg, test=True)
    assert [o.kind for o in outs] == ["new"]
    first = outs[0]
    assert first.package_name.startswith("Test-CRS-2026-") and first.package is not None
    assert packaging.unpackage(first.package, key) == first.result.xml.encode("utf-8")
    view = workflow.registry_view(reg)
    assert view["pending"] == [first.result.message_ref_id] and view["counts"] == {"built": 1}

    # 2. before the portal answered: nothing can be corrected
    intent = workflow.plan(msg, "3.0", reg, test=True)
    assert intent.blocked == ["A1", "A2"] and "record the portal result" in intent.problems[0]
    with pytest.raises(workflow.WorkflowError, match="portal result"):
        workflow.build(msg, "3.0", reg, test=True)

    # 3. accepted; the same workbook means nothing to send
    _accept(reg, first.result.message_ref_id)
    intent = workflow.plan(msg, "3.0", reg, test=True)
    assert intent.unchanged == ["A1", "A2"] and intent.messages == []
    with pytest.raises(workflow.WorkflowError, match="unchanged"):
        workflow.build(msg, "3.0", reg, test=True)

    # 4. one changed row, one cancelled, one new account -> correction + new message
    changed = msg.model_copy(deep=True)
    changed.accounts[0].balance = changed.accounts[0].balance + 1
    third = changed.accounts[0].model_copy(
        update={"key": "A3", "doc_ref_id": None, "account_number": "CH5604835012345678009"}
    )
    changed.accounts.append(third)
    intent = workflow.plan(changed, "3.0", reg, test=True, cancel=["A2"])
    assert intent.changed == ["A1"] and intent.deletions == ["A2"] and intent.new_keys == ["A3"]
    assert intent.messages == ["correction", "new"]
    outs = workflow.build(changed, "3.0", reg, test=True, cancel=["A2"])
    assert [o.kind for o in outs] == ["correction", "new"]
    corr, new = outs
    assert _indics(corr.result.xml) == [
        "OECD10",
        "OECD12",
        "OECD13",
    ]  # FI resent, A1 corrected, A2 deleted
    assert _indics(new.result.xml) == ["OECD10", "OECD11"]
    assert set(corr.result.doc_ref_ids) == {"A1", "A2"} and list(new.result.doc_ref_ids) == ["A3"]
    assert corr.package is not None and new.package is not None

    # 5. after acceptance, the cancelled account comes back as a new record (98103)
    _accept(reg, corr.result.message_ref_id)
    _accept(reg, new.result.message_ref_id)
    intent = workflow.plan(msg, "3.0", reg, test=True)
    assert intent.new_keys == ["A2"] and intent.changed == ["A1"]


def test_build_from_workbook_refuses_errors(tmp_path, reg):
    path = tmp_path / "book.xlsx"
    template.write_message(sample_message(), path)
    report, outs = workflow.build_from_workbook(path, "3.0", reg, test=True)
    assert report.ok and outs[0].kind == "new" and outs[0].package is None  # no key: XML only
    assert outs[0].as_dict()["encrypted"] is False and outs[0].xml_name.endswith(".xml")
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb["Accounts"]
    ws.cell(row=2, column=[c.value for c in ws[1]].index("last_name") + 1).value = "Mü#ller"
    wb.save(path)
    with pytest.raises(workflow.WorkflowError, match="errors"):
        workflow.build_from_workbook(path, "3.0", reg, test=True)


def test_record_outcome_from_text_and_xml(reg):
    outs = workflow.build(sample_message(), "3.0", reg, test=True)
    ref = outs[0].result.message_ref_id
    with pytest.raises(workflow.WorkflowError, match="MessageRefId"):
        workflow.record_outcome(reg, "Status: rejected 50005")
    with pytest.raises(workflow.WorkflowError, match="not in this registry"):
        workflow.record_outcome(reg, "rejected 50005", message_ref_id="CH2026CHnope")
    view = workflow.record_outcome(  # the page passes the message chosen from the pending list
        reg, f"abgelehnt: 50005 {outs[0].result.doc_ref_ids['A1']}", message_ref_id=ref, lang="it"
    )
    assert view["status"] == "rejected" and view["accepted"] is False
    assert view["findings"][0]["code"] == "50005"
    assert view["findings"][0]["title"] == "Carattere non ammesso"
    assert view["findings"][0]["should_have_been_caught"] is True
    assert workflow.registry_view(reg)["counts"] == {"rejected": 1}


def test_portal_status_fehler_is_not_a_verdict(reg):
    """Benutzeranleitung: «Der Status «Fehler» erscheint, falls ein unbekanntes Problem die
    Verarbeitung verhinderte. In diesem Fall muss die Datei noch einmal hochgeladen werden.»"""
    outs = workflow.build(sample_message(), "3.0", reg, test=True)
    ref = outs[0].result.message_ref_id
    outcome = workflow.parse_outcome("Meldungsübersicht\nStatus: «Fehler»\n")
    assert outcome.portal_error and outcome.accepted is None and not outcome.findings
    view = workflow.record_outcome(reg, "Status: Fehler", message_ref_id=ref, lang="de")
    assert view["status"] == "submitted" and view["accepted"] is None
    assert (
        view["note"]
        == "Status «Fehler»: unbekanntes Problem beim Portal; die Datei noch einmal hochladen"
    )
    assert workflow.registry_view(reg)["counts"] == {"submitted": 1}
    # the message stays open: it can still be accepted afterwards, and "Fehlerbericht" is a rejection
    assert not workflow.parse_outcome("Fehlerbericht: 50005").portal_error
    _accept(reg, ref)
    assert workflow.registry_view(reg)["counts"] == {"accepted": 1}


def test_header_variant_and_key_formats(reg, key_pair, tmp_path):
    """Both 3.0 headers build, register and encrypt; the ESTV key is accepted as PEM or DER,
    key or certificate, and is stored as PEM."""
    import datetime as dt

    from cryptography import x509
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.x509.oid import NameOID

    from aeoi.crs import build as builder
    from aeoi.estv import packaging

    key, pem = key_pair
    outs = workflow.build(sample_message(), "3.0", reg, test=True, header="wegleitung")
    assert outs[0].result.header == "wegleitung" and builder.is_wegleitung_header(
        outs[0].result.xml
    )
    assert outs[0].as_dict()["header"] == "wegleitung"
    outs2 = workflow.build(sample_message(year=2025), "3.0", reg, test=True)
    assert outs2[0].result.header == "oecd" and not builder.is_wegleitung_header(
        outs2[0].result.xml
    )

    der_key = key.public_key().public_bytes(
        serialization.Encoding.DER, serialization.PublicFormat.SubjectPublicKeyInfo
    )
    name = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "ESTV test")])
    cert = (
        x509.CertificateBuilder()
        .subject_name(name)
        .issuer_name(name)
        .public_key(key.public_key())
        .serial_number(1)
        .not_valid_before(dt.datetime(2026, 1, 1, tzinfo=dt.UTC))
        .not_valid_after(dt.datetime(2030, 1, 1, tzinfo=dt.UTC))
        .sign(key, hashes.SHA256())
    )
    forms = {
        "pem key": pem,
        "der key": der_key,
        "pem cert": cert.public_bytes(serialization.Encoding.PEM),
        "der cert": cert.public_bytes(serialization.Encoding.DER),
    }
    digests = {n: workflow.remember_public_key(reg, data) for n, data in forms.items()}
    assert len(set(digests.values())) == 1, digests
    assert reg.get_setting(workflow.PUBLIC_KEY_SETTING).startswith("-----BEGIN PUBLIC KEY-----")
    with pytest.raises(packaging.PackagingError, match="PEM or DER"):
        packaging.load_public_key(b"not a key")


def test_registry_remembers_its_institution(reg, tmp_path):
    """A build with a registry records which FI the file belongs to (the Mandantenübersicht
    pairs workbooks and registries by it); a fresh registry knows nothing; a restored one learns
    it from the sent XML files."""
    assert reg.fi() is None and workflow.registry_view(reg)["fi"] is None
    msg = sample_message()
    outs = workflow.build(msg, "3.0", reg, test=True)
    assert reg.fi() == {"estv_id": msg.reporting_fi.estv_id, "name": msg.reporting_fi.name}
    assert workflow.registry_view(reg)["fi"]["name"] == msg.reporting_fi.name
    xml = tmp_path / "sent.xml"
    xml.write_text(outs[0].result.xml, encoding="utf-8")
    fresh = Registry(tmp_path / "fresh.sqlite")
    workflow.restore_registry(fresh, [xml], status="accepted")
    assert fresh.fi() == reg.fi()
    fresh.close()
